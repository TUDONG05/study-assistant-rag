from __future__ import annotations

from io import BytesIO

from docx import Document

from src.ingestion.models import UploadPayload
from src.ingestion.parsers import parse_document
from src.ingestion.validation import validate_upload


def _docx_with_interleaved_table() -> bytes:
    document = Document()
    document.add_heading("Phần một", level=1)
    document.add_paragraph("Nội dung đầu tiên đủ dài để được parser chấp nhận.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Khái niệm"
    table.cell(0, 1).text = "Giải thích"
    document.add_heading("Phần hai", level=1)
    document.add_paragraph("Nội dung thứ hai nằm sau bảng trong tài liệu.")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_parser_preserves_table_order_and_section() -> None:
    upload = validate_upload(
        UploadPayload(
            file_name="bai-hoc.docx",
            data=_docx_with_interleaved_table(),
            mime_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
        ),
        max_upload_bytes=2_000_000,
        max_zip_entries=1_000,
        max_zip_uncompressed_bytes=10_000_000,
    )

    parsed = parse_document(upload, max_pages=20, max_extracted_chars=10_000)

    assert [(block.source, block.section) for block in parsed.blocks] == [
        ("Đoạn 1", "Phần một"),
        ("Bảng 1", "Phần một"),
        ("Đoạn 2", "Phần hai"),
    ]

