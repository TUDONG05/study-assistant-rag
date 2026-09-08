from __future__ import annotations

from io import BytesIO

from docx import Document
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def make_docx_bytes(
    *,
    first_text: str = "Nội dung đầu tiên đủ dài để được parser chấp nhận.",
    include_table: bool = False,
) -> bytes:
    document = Document()
    document.add_heading("Phần một", level=1)
    document.add_paragraph(first_text)
    if include_table:
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Khái niệm"
        table.cell(0, 1).text = "Giải thích"
        document.add_heading("Phần hai", level=1)
        document.add_paragraph("Nội dung thứ hai nằm sau bảng trong tài liệu.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Bài giảng RAG"
    slide.placeholders[1].text = "Nội dung slide đủ dài để trích xuất và lập chỉ mục."
    table = slide.shapes.add_table(1, 2, 0, 2_000_000, 4_000_000, 800_000).table
    table.cell(0, 0).text = "Dense"
    table.cell(0, 1).text = "Vector"
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_pdf_bytes(text: str = "A sufficiently long lesson about retrieval.") -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)  # noqa: SLF001 - pypdf has no public font API
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_reference})}
    )
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(content)  # noqa: SLF001
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()
