"""Parsers that preserve citation boundaries and source metadata."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from src.ingestion.models import DocumentKind, ParsedBlock, ParsedDocument, ValidatedUpload
from src.ingestion.validation import IngestionError


def parse_document(
    upload: ValidatedUpload,
    *,
    max_pages: int,
    max_extracted_chars: int,
) -> ParsedDocument:
    try:
        if upload.kind is DocumentKind.PDF:
            blocks = _parse_pdf(
                upload,
                max_pages=max_pages,
                max_extracted_chars=max_extracted_chars,
            )
        elif upload.kind is DocumentKind.DOCX:
            blocks = _parse_docx(upload)
        else:
            blocks = _parse_pptx(upload, max_slides=max_pages)
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError(f"{upload.file_name}: không thể đọc nội dung tài liệu.") from exc

    extracted_chars = sum(len(block.text) for block in blocks)
    if extracted_chars < 20:
        detail = " PDF có thể là bản scan cần OCR." if upload.kind is DocumentKind.PDF else ""
        raise IngestionError(f"{upload.file_name}: không tìm thấy đủ văn bản.{detail}")
    if extracted_chars > max_extracted_chars:
        raise IngestionError(
            f"{upload.file_name}: văn bản trích xuất vượt giới hạn {max_extracted_chars:,} ký tự."
        )
    return ParsedDocument(
        title=upload.title,
        kind=upload.kind,
        blocks=tuple(blocks),
        extracted_chars=extracted_chars,
    )


def _parse_pdf(
    upload: ValidatedUpload,
    *,
    max_pages: int,
    max_extracted_chars: int,
) -> list[ParsedBlock]:
    try:
        reader = PdfReader(BytesIO(upload.data), strict=False)
    except PdfReadError as exc:
        raise IngestionError(f"{upload.file_name}: PDF bị hỏng hoặc không đọc được.") from exc
    if reader.is_encrypted:
        raise IngestionError(f"{upload.file_name}: không hỗ trợ PDF có mật khẩu.")
    if len(reader.pages) > max_pages:
        raise IngestionError(f"{upload.file_name}: vượt giới hạn {max_pages} trang.")

    blocks: list[ParsedBlock] = []
    extracted_chars = 0
    for index, page in enumerate(reader.pages):
        text = _clean_text(page.extract_text() or "")
        if text:
            extracted_chars += len(text)
            if extracted_chars > max_extracted_chars:
                raise IngestionError(
                    f"{upload.file_name}: văn bản trích xuất vượt giới hạn "
                    f"{max_extracted_chars:,} ký tự."
                )
            page_number = index + 1
            blocks.append(
                ParsedBlock(
                    block_index=index,
                    text=text,
                    source=f"Trang {page_number}",
                    page_number=page_number,
                )
            )
    return blocks


def _parse_docx(upload: ValidatedUpload) -> list[ParsedBlock]:
    document = DocxDocument(BytesIO(upload.data))
    blocks: list[ParsedBlock] = []
    section: str | None = None
    block_index = 0
    paragraph_index = 0
    table_index = 0

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = _clean_text(item.text)
            if not text:
                continue
            style_name = str(getattr(item.style, "name", ""))
            if style_name.casefold().startswith("heading"):
                section = text
                continue
            paragraph_index += 1
            blocks.append(
                ParsedBlock(
                    block_index=block_index,
                    text=text,
                    source=f"Đoạn {paragraph_index}",
                    section=section,
                )
            )
            block_index += 1
            continue

        if isinstance(item, Table):
            rows = []
            for row in item.rows:
                cells = [_clean_text(cell.text) for cell in row.cells]
                line = " | ".join(cell for cell in cells if cell)
                if line:
                    rows.append(line)
            if rows:
                table_index += 1
                blocks.append(
                    ParsedBlock(
                        block_index=block_index,
                        text="\n".join(rows),
                        source=f"Bảng {table_index}",
                        section=section,
                    )
                )
                block_index += 1
    return blocks


def _parse_pptx(upload: ValidatedUpload, *, max_slides: int) -> list[ParsedBlock]:
    presentation = Presentation(BytesIO(upload.data))
    if len(presentation.slides) > max_slides:
        raise IngestionError(f"{upload.file_name}: vượt giới hạn {max_slides} slide.")

    blocks: list[ParsedBlock] = []
    for index, slide in enumerate(presentation.slides):
        texts: list[str] = []
        title = _clean_text(slide.shapes.title.text) if slide.shapes.title else ""
        for shape in slide.shapes:
            text = _pptx_shape_text(shape)
            if text and text not in texts:
                texts.append(text)
        combined = "\n\n".join(texts)
        if combined:
            slide_number = index + 1
            blocks.append(
                ParsedBlock(
                    block_index=index,
                    text=combined,
                    source=f"Slide {slide_number}",
                    section=title or None,
                    slide_number=slide_number,
                )
            )
    return blocks


def _pptx_shape_text(shape: Any) -> str:
    if getattr(shape, "has_text_frame", False):
        return _clean_text(str(getattr(shape, "text", "")))
    if not getattr(shape, "has_table", False):
        return ""

    rows = []
    for row in shape.table.rows:
        cells = [_clean_text(cell.text) for cell in row.cells]
        line = " | ".join(cell for cell in cells if cell)
        if line:
            rows.append(line)
    return "\n".join(rows)


def _clean_text(value: Any) -> str:
    text = str(value).replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[\t ]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()
